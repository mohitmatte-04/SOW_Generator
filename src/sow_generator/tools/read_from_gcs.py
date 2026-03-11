"""Tool to read content from Google Cloud Storage."""

import json
import logging
from pathlib import Path
from typing import Any

from docx import Document

from sow_generator.utils.gcs_utils import download_blob_to_tempfile

logger = logging.getLogger(__name__)


def _extract_text_from_docx(local_path: Path) -> str:
    """Extract human-readable text from a .docx file.

    Returns a string representation including paragraph text, table
    content, and placeholder tags so the LLM can understand the
    template structure.
    """
    doc = Document(str(local_path))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_idx, table in enumerate(doc.tables):
        parts.append(f"\n[Table {table_idx + 1}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


async def read_from_gcs(gcs_uri: str) -> dict[str, Any]:
    """Read a file from GCS and return its content.

    Supports the following file types:
    - ``.json`` → parsed as JSON and returned as a dictionary.
    - ``.docx`` → text content is extracted (paragraphs + tables) and
      returned as a string so the LLM can inspect the template structure.
    - Other text files → raw text content is returned.

    Args:
        gcs_uri: Full ``gs://…`` URI to the object.

    Returns:
        A dictionary with ``status``, ``content_type``, and ``content``.
    """
    try:
        local_path = download_blob_to_tempfile(gcs_uri)

        suffix = local_path.suffix.lower()

        if suffix == ".json":
            content = json.loads(local_path.read_text(encoding="utf-8"))
            content_type = "json"
        elif suffix == ".docx":
            content = _extract_text_from_docx(local_path)
            content_type = "docx_text"
        else:
            content = local_path.read_text(encoding="utf-8")
            content_type = "text"

        # Clean up temp file
        local_path.unlink(missing_ok=True)

        logger.info("Read %s from %s (%s)", content_type, gcs_uri, suffix)
        return {
            "status": "success",
            "content_type": content_type,
            "content": content,
        }

    except Exception as e:
        logger.error("Failed to read from GCS: %s", e, exc_info=True)
        return {"status": "error", "error": str(e)}
