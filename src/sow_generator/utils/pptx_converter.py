"""PPTX to PDF converter utilities.

Provides conversion from PowerPoint to PDF format for downstream
multimodal processing with Gemini. Uses LibreOffice headless as the
primary converter with a python-pptx text-extraction fallback.
"""

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

logger = logging.getLogger(__name__)


class ConversionError(RuntimeError):
    """Raised when PPTX-to-PDF conversion fails."""


def convert_pptx_to_pdf(pptx_path: Path) -> Path:
    """Convert a PPTX file to PDF using LibreOffice headless.

    Args:
        pptx_path: Path to the source ``.pptx`` file.

    Returns:
        Path to the generated PDF file.

    Raises:
        ConversionError: If LibreOffice is not installed or the conversion
            fails.
        FileNotFoundError: If the source PPTX file does not exist.
    """
    if not pptx_path.exists():
        msg = f"PPTX file not found: {pptx_path}"
        raise FileNotFoundError(msg)

    libreoffice_bin = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice_bin is None:
        msg = (
            "LibreOffice is not installed. Install it with "
            "'apt-get install -y libreoffice' or use extract_text_from_pptx() "
            "as a fallback."
        )
        raise ConversionError(msg)

    outdir = Path(tempfile.mkdtemp(prefix="sow_pptx_"))
    logger.info("Converting %s → PDF in %s", pptx_path, outdir)

    result = subprocess.run(  # noqa: S603
        [
            libreoffice_bin,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(pptx_path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        msg = f"LibreOffice conversion failed: {result.stderr}"
        raise ConversionError(msg)

    pdf_path = outdir / pptx_path.with_suffix(".pdf").name
    if not pdf_path.exists():
        msg = f"Expected PDF not found at {pdf_path} after conversion"
        raise ConversionError(msg)

    logger.info("PDF created: %s", pdf_path)
    return pdf_path


def extract_text_from_pptx(pptx_path: Path) -> str:
    """Extract text content from a PPTX file slide-by-slide.

    This is a fallback for environments where LibreOffice is not available.
    It preserves slide ordering and labels each slide.

    Args:
        pptx_path: Path to the ``.pptx`` file.

    Returns:
        A formatted string with text content grouped by slide number.

    Raises:
        FileNotFoundError: If the PPTX file does not exist.
    """
    if not pptx_path.exists():
        msg = f"PPTX file not found: {pptx_path}"
        raise FileNotFoundError(msg)

    prs = Presentation(str(pptx_path))
    slides_text: list[str] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = [f"--- Slide {slide_num} ---"]

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_lines.append(text)

            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells
                    )
                    if row_text.strip():
                        slide_lines.append(row_text)

        slides_text.append("\n".join(slide_lines))

    return "\n\n".join(slides_text)
