import json
from docx import Document
from docx.shared import Pt, Inches

def create_sow(json_path, output_path):
    with open(json_path, 'r') as f:
        data = json.load(f)
        
    doc = Document()
    
    # Title
    title = doc.add_heading('Statement of Work (SOW)', 0)
    title.alignment = 1 # Center
    
    doc.add_paragraph('Prepared By: Onix Networking Corp.')
    doc.add_paragraph('Customer: Bell Canada')
    doc.add_paragraph('')
    
    def add_section(title, items):
        doc.add_heading(title, level=1)
        if isinstance(items, list):
            for item in items:
                doc.add_paragraph(item, style='List Bullet')
        else:
            doc.add_paragraph(str(items))
        doc.add_paragraph('')

    add_section('1.0 Solution Overview', data.get('solution_overview', []))
    add_section('2.0 Scope', data.get('scope', []))
    add_section('3.0 Deliverables', data.get('deliverables', []))
    add_section('4.0 Success Criteria', data.get('success_criteria', []))
    add_section('5.0 Technical Assumptions', data.get('technical_assumptions', []))
    add_section('6.0 Out of Scope', data.get('out_of_scope', []))

    doc.save(output_path)

if __name__ == "__main__":
    create_sow('enriched-content-Bell.json', 'Final-SOW-Bell.docx')
