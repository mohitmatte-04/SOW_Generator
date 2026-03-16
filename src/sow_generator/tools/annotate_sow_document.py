"""Tool to annotate SOW document with validation feedback."""

import logging
import io
from typing import Any, Dict, List
from google.cloud import storage
from docx import Document
from docx.shared import Pt, RGBColor
from ..config import VALIDATION_QUALITY_THRESHOLD

logger = logging.getLogger(__name__)

async def annotate_sow_document(
    gcs_uri: str,
    section_feedback: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """
    Annotates an existing SOW DOCX with manual update notices for low-scoring sections.
    Supports searching in both main paragraphs and tables.

    Args:
        gcs_uri: GCS URI of the target DOCX file.
        section_feedback: Feedback from validation agent containing scores and sections.

    Returns:
        Status and GCS URI of the annotated document.
    """
    if not section_feedback:
        logger.info("No section feedback provided for annotation.")
        return {
            "status": "success",
            "data": {
                "gcs_uri": gcs_uri,
                "annotations_added": 0
            }
        }

    try:
        storage_client = storage.Client()

        # 1. Download document from GCS
        uri_str = str(gcs_uri)
        if not uri_str.startswith("gs://"):
            raise ValueError(f"Invalid GCS URI: {gcs_uri}")
            
        parts = uri_str.replace("gs://", "").split("/")
        bucket_name = parts[0]
        blob_path = "/".join(parts[1:])
        
        bucket = storage_client.bucket(bucket_name)
        blob = bucket.blob(blob_path)
        
        doc_stream = io.BytesIO()
        blob.download_to_file(doc_stream)
        doc_stream.seek(0)

        # 2. Load DOCX
        document = Document(doc_stream)

        # 3. Apply annotations
        annot_count = 0
        threshold = VALIDATION_QUALITY_THRESHOLD
        
        # Build list of all paragraphs (main + tables)
        all_paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_paragraphs.extend(cell.paragraphs)

        for fb in section_feedback:
            if not isinstance(fb, dict):
                continue
                
            section_name = str(fb.get("section", ""))
            score_val = fb.get("score")
            score = int(score_val) if score_val is not None else 100
            
            # Use synchronized threshold
            if score < threshold:
                target_para = None
                original_text = fb.get("original_text", "")
                
                # Strategy 1: Find paragraph by original text snippet
                if original_text:
                    text_str = str(original_text)
                    snippet = text_str[0:100].strip() if len(text_str) > 100 else text_str.strip()
                    
                    for para in all_paragraphs:
                        if snippet in para.text:
                            target_para = para
                            break
                
                # Strategy 2: Fallback to section name hint
                if not target_para and section_name:
                    hint = section_name.replace("<<", "").replace(">>", "").replace("_", " ")
                    for para in all_paragraphs:
                        if hint.lower() in para.text.lower():
                            target_para = para
                            break
                
                if target_para:
                    # Insert attention message
                    # Handling insertion: if target is in a table, we add to the cell
                    parent = target_para._element.getparent()
                    
                    # Create a new paragraph element using the same parent structure
                    new_para = target_para.insert_paragraph_before("") # Workaround to get correct parent
                    # Move it to after
                    target_para._element.addnext(new_para._element)
                    
                    new_para.paragraph_format.space_before = Pt(6)
                    new_para.paragraph_format.space_after = Pt(6)
                    
                    run_header = new_para.add_run("⚠️ ATTENTION:")
                    run_header.bold = True
                    run_header.font.color.rgb = RGBColor(0xFF, 0x66, 0x00)
                    run_header.font.size = Pt(10)
                    
                    new_para.add_run(" Notice: This section needs additional inputs or needs to be updated manually.")
                    run_msg = new_para.runs[-1]
                    run_msg.bold = True
                    run_msg.font.color.rgb = RGBColor(0xFF, 0x66, 0x00)
                    run_msg.font.size = Pt(9)
                    
                    annot_count += 1

        # 4. Save modified doc
        output_stream = io.BytesIO()
        document.save(output_stream)
        output_stream.seek(0)

        # 5. Overwrite the file on GCS
        blob.upload_from_file(
            output_stream,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        logger.info(f"Annotated {annot_count} sections in {gcs_uri} (Threshold: {threshold})")

        return {
            "status": "success",
            "data": {
                "gcs_uri": gcs_uri,
                "annotations_added": annot_count,
                "threshold_used": threshold
            }
        }

    except Exception as e:
        logger.error(f"Failed to annotate SOW document: {e}", exc_info=True)
        return {"status": "error", "error": str(e)}
