from docx import Document

doc = Document('gws-app/template.docx')
for para in doc.paragraphs:
    if "<<" in para.text:
        print("PARA:", para.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if "<<" in para.text:
                    print("CELL:", para.text)
